"""Local reference ingestion. Parsing runs in a bounded, isolated subprocess."""
from __future__ import annotations

from contextlib import closing
import io
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import zipfile

from .util import atomic_json, atomic_write, digest, read_json, validate_id

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TEXT_BYTES = 2 * 1024 * 1024
MAX_REFERENCES = 30
EXTENSIONS = {".md", ".txt", ".pdf", ".pptx", ".ppt", ".docx", ".png", ".jpg", ".jpeg", ".webp"}


def ocr_image(blob, warnings):
    from PIL import Image, ImageOps
    with Image.open(io.BytesIO(blob)) as source:
        if source.width * source.height > 25_000_000:
            raise ValueError("图片超过 2500 万像素，请缩小后上传")
        image = ImageOps.exif_transpose(source).convert("RGB")
        image.thumbnail((3200, 3200))
        with tempfile.TemporaryDirectory(prefix="rsi-ocr-") as tmp:
            path = Path(tmp) / "image.png"
            image.save(path)
            if sys.platform == "darwin":
                try:
                    import Vision
                    from Foundation import NSURL
                    request = Vision.VNRecognizeTextRequest.alloc().init()
                    request.setRecognitionLevel_(Vision.VNRequestTextRecognitionLevelAccurate)
                    request.setRecognitionLanguages_(["zh-Hans", "zh-Hant", "en-US"])
                    request.setUsesLanguageCorrection_(True)
                    handler = Vision.VNImageRequestHandler.alloc().initWithURL_options_(NSURL.fileURLWithPath_(str(path)), {})
                    ok, error = handler.performRequests_error_([request], None)
                    if ok:
                        return "\n".join(item.topCandidates_(1)[0].string() for item in request.results() if item.topCandidates_(1)).strip()
                except ImportError:
                    warnings.append("本机缺少 macOS Vision 依赖，请重新安装项目依赖。")
            if shutil.which("tesseract"):
                languages = subprocess.run(["tesseract", "--list-langs"], capture_output=True, text=True, timeout=10).stdout.splitlines()
                chosen = [lang for lang in ("chi_sim", "chi_tra", "eng") if lang in languages]
                if chosen:
                    if "chi_sim" not in chosen:
                        warnings.append("OCR 未安装简体中文语言包，中文图片可能无法正确识别；请核对或补充文字。")
                    result = subprocess.run(["tesseract", str(path), "stdout", "-l", "+".join(chosen)],
                                            capture_output=True, text=True, timeout=40)
                    if result.returncode == 0:
                        return result.stdout.strip()
            warnings.append("图像文字未识别：macOS 需可用的 Vision 依赖，其他系统需安装 Tesseract 及中文语言包。")
            return ""


def check_archive(blob):
    with zipfile.ZipFile(io.BytesIO(blob)) as archive:
        if len(archive.infolist()) > 5000 or sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError("文档解压后过大，请拆分文件")


def extract(blob, suffix):
    warnings = []
    parts = []
    if suffix in {".md", ".txt"}:
        parts.append(blob.decode("utf-8-sig"))
    elif suffix == ".pdf":
        import pypdfium2 as pdfium
        warnings.append("PDF 按页提取文字；扫描页使用 OCR。图表关系、版式和混合页面中的图片信息需人工核对。")
        with pdfium.PdfDocument(blob) as pdf:
            if len(pdf) > 100:
                raise ValueError("PDF 超过 100 页，请拆分后上传")
            for i in range(len(pdf)):
                with closing(pdf[i]) as page:
                    with closing(page.get_textpage()) as textpage:
                        text = textpage.get_text_bounded().strip()
                    if not text:
                        scale = min(2.0, 3200 / max(page.get_size()))
                        with closing(page.render(scale=scale)) as bitmap:
                            buffer = io.BytesIO()
                            bitmap.to_pil().save(buffer, format="PNG")
                            text = ocr_image(buffer.getvalue(), warnings)
                    if not text:
                        warnings.append(f"第 {i + 1} 页没有可识别文字，请补充该页信息。")
                    parts.append(f"[第 {i + 1} 页]\n{text}")
    elif suffix == ".ppt":
        executable = shutil.which("soffice") or shutil.which("libreoffice")
        if not executable:
            raise ValueError("旧版 .ppt 需安装 LibreOffice；也可用 PowerPoint 另存为 .pptx 后上传")
        with tempfile.TemporaryDirectory(prefix="rsi-ppt-") as tmp:
            folder = Path(tmp)
            source = folder / "reference.ppt"
            source.write_bytes(blob)
            result = subprocess.run([executable, "-env:UserInstallation=" + (folder / "profile").as_uri(),
                                     "--headless", "--convert-to", "pptx", "--outdir", tmp, str(source)],
                                    capture_output=True, timeout=60)
            target = folder / "reference.pptx"
            if result.returncode or not target.exists():
                raise ValueError("旧版 PPT 转换失败，请另存为 .pptx 后上传")
            return extract(target.read_bytes(), ".pptx")
    elif suffix == ".pptx":
        from pptx import Presentation
        check_archive(blob)
        deck = Presentation(io.BytesIO(blob))
        if len(deck.slides) > 100:
            raise ValueError("幻灯片超过 100 页，请拆分后上传")
        warnings.append("已提取幻灯片文字、表格、图表数据、备注及可识别图片文字；图形关系和图像含义需补充说明。")

        def shapes_text(shapes):
            texts = []
            for shape in shapes:
                if hasattr(shape, "shapes"):
                    texts.extend(shapes_text(shape.shapes))
                if shape.has_text_frame:
                    texts.append(shape.text)
                if shape.has_table:
                    texts.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
                if shape.has_chart:
                    chart = shape.chart
                    if chart.has_title and chart.chart_title.has_text_frame:
                        texts.append(chart.chart_title.text_frame.text)
                    try:
                        for plot in chart.plots:
                            texts.append("图表类别：" + " / ".join(str(c.label) for c in plot.categories))
                            for series in plot.series:
                                texts.append(str(series.name) + "：" + " / ".join(str(v) for v in series.values))
                    except (AttributeError, ValueError, TypeError):
                        warnings.append("有图表序列无法完整提取，请补充原始数据及图表说明。")
                if hasattr(shape, "image"):
                    try:
                        recognized = ocr_image(shape.image.blob, warnings)
                        if recognized:
                            texts.append("图片文字：" + recognized)
                    except (OSError, ValueError):
                        warnings.append("有嵌入图片无法识别，请核对原幻灯片。")
            return texts

        for i, slide in enumerate(deck.slides, 1):
            text = shapes_text(slide.shapes)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame and slide.notes_slide.notes_text_frame.text.strip():
                text.append("演讲备注：" + slide.notes_slide.notes_text_frame.text)
            parts.append(f"[第 {i} 张幻灯片]\n" + "\n".join(text))
    elif suffix == ".docx":
        from docx import Document
        check_archive(blob)
        doc = Document(io.BytesIO(blob))
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        for block in doc.iter_inner_content():
            if isinstance(block, Paragraph):
                parts.append(block.text)
            elif isinstance(block, Table):
                parts.extend(" | ".join(c.text for c in row.cells) for row in block.rows)
        for relation in doc.part.rels.values():
            if not relation.is_external and relation.reltype.endswith("/image"):
                try:
                    recognized = ocr_image(relation.target_part.blob, warnings)
                    if recognized:
                        parts.append("图片文字：" + recognized)
                except (OSError, ValueError):
                    warnings.append("Word 中有图片无法识别，请核对原文件。")
        warnings.append("已提取 Word 正文、表格和图片文字；页眉页脚、批注、文本框及修订标记不作为正文。")
    elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        parts.append(ocr_image(blob, warnings))
        warnings.append("图片通过本地 OCR 转成文字；无法据此判断图像、布局或曲线含义，请补充说明。")
    else:
        raise ValueError("不支持此文件格式")
    text = "\n\n".join(parts).strip()
    # Page labels do not count as usable source content.
    substantive = "\n".join(line for line in text.splitlines() if not line.startswith("[第 "))
    if not substantive.strip():
        raise ValueError("未识别到可用文字。请提供清晰扫描件、可复制文本或补充文字说明。" + " ".join(dict.fromkeys(warnings)))
    if len(text.encode()) > MAX_TEXT_BYTES:
        raise ValueError("提取文字超过 2 MB，请拆分资料")
    return {"text": text, "warnings": list(dict.fromkeys(warnings)), "parser_version": 1}


class References:
    def __init__(self, state):
        self.root = Path(state) / "references"
        self.root.mkdir(exist_ok=True)

    def get(self, reference_id):
        row = read_json(self.root / validate_id(reference_id) / "reference.json")
        if digest(row["text"]) != row["text_hash"]:
            raise ValueError("参考资料解析内容发生变化，请重新上传")
        return row

    def upload(self, name, blob):
        if not name or len(name) > 240 or "/" in name or "\\" in name or any(ord(c) < 32 for c in name):
            raise ValueError("文件名无效")
        suffix = Path(name).suffix.lower()
        if suffix not in EXTENSIONS:
            raise ValueError("支持 PDF、PPT/PPTX、DOCX、MD、TXT、PNG、JPG、WEBP")
        if not blob or len(blob) > MAX_FILE_BYTES:
            raise ValueError("每份参考资料须为 1 字节至 20 MB")
        reference_id = "ref_" + digest(name.encode() + b"\0" + blob)[:32]
        folder = self.root / reference_id
        if (folder / "reference.json").exists():
            return self.get(reference_id)
        with tempfile.TemporaryDirectory(prefix="rsi-parse-") as tmp:
            source, output = Path(tmp) / ("source" + suffix), Path(tmp) / "parsed.json"
            source.write_bytes(blob)
            try:
                result = subprocess.run([sys.executable, "-m", "writing_memory.references", str(source), str(output)],
                                        capture_output=True, text=True, timeout=180)
            except subprocess.TimeoutExpired:
                raise ValueError("解析超过 3 分钟，未导入。请拆分资料后重试。") from None
            if not output.exists():
                raise ValueError("文件解析失败，请检查文件是否损坏或加密，并尝试另存后重传")
            parsed = read_json(output)
            if result.returncode or "error" in parsed:
                raise ValueError(parsed.get("error", "文件解析失败"))
        row = {**parsed, "id": reference_id, "name": name, "size": len(blob), "source_hash": digest(blob),
               "text_hash": digest(parsed["text"]), "text_bytes": len(parsed["text"].encode())}
        atomic_write(folder / ("original" + suffix), blob)
        atomic_json(folder / "reference.json", row)
        return row

    def select(self, ids):
        if len(ids) > MAX_REFERENCES or len(set(ids)) != len(ids):
            raise ValueError("每篇最多选择 30 份不同参考资料")
        return [self.get(i) for i in ids]


if __name__ == "__main__":
    source, output = map(Path, sys.argv[1:])
    try:
        atomic_json(output, extract(source.read_bytes(), source.suffix))
    except Exception as exc:
        atomic_json(output, {"error": "参考资料解析失败：" + str(exc)[:500]})
        sys.exit(1)
